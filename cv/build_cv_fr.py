from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

FONT = "Calibri"

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
rpr = normal.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = rpr.makeelement(qn('w:rFonts'), {})
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), FONT)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)

def add_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pPr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '6', qn('w:space'): '1', qn('w:color'): '000000'
    })
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_entry_title(title, meta):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(10.5)
    if meta:
        p.add_run("\t")
        r2 = p.add_run(meta)
        r2.italic = True
        r2.font.size = Pt(10)
        from docx.enum.text import WD_TAB_ALIGNMENT
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    return p

def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    return p

def add_plain(text, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    return p

# ---------- EN-TETE ----------
name_p = doc.add_paragraph()
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
name_p.paragraph_format.space_after = Pt(2)
r = name_p.add_run("DEIDINE SIDINA")
r.bold = True
r.font.size = Pt(20)

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_after = Pt(4)
r = title_p.add_run("Ingénieur Logiciel")
r.font.size = Pt(12.5)
r.italic = True

contact_p = doc.add_paragraph()
contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_p.paragraph_format.space_after = Pt(2)
r = contact_p.add_run(
    "cheigeurdeidine@gmail.com  |  +222 49619609  |  linkedin.com/in/deidine  |  github.com/deidine"
)
r.font.size = Pt(10)

# ---------- PROFIL ----------
add_heading("Profil")
add_plain(
    "Ingénieur logiciel avec une expérience dans la livraison d'applications web, mobiles et "
    "desktop full-stack pour des projets dans l'immobilier, l'e-commerce, le secteur public et "
    "la fintech. Maîtrise de React/Next.js, React Native, Spring Boot et Python, avec une "
    "expérience pratique de l'intégration IA (OpenAI SDK, RAG), du déploiement cloud (AWS, GCP) "
    "et de la conception de bases de données SQL et NoSQL.",
    space_after=2,
)

# ---------- COMPETENCES ----------
add_heading("Compétences")
skills = [
    ("Langages", "JavaScript, TypeScript, Python, Java, C, C++, PHP, Dart, SQL"),
    ("Frontend & Mobile", "React, Next.js, Angular, Vue.js, React Native, Electron, HTML, CSS, Tailwind CSS, Bootstrap"),
    ("Backend", "Spring Boot, Hibernate/JPA, Laravel, Django, REST API, SOAP"),
    ("Bases de données", "MySQL, PostgreSQL, MongoDB, Firebase, Supabase, SQL Server, NoSQL"),
    ("IA & Données", "Machine Learning, Deep Learning, NLP, OpenAI SDK, Retrieval-Augmented Generation (RAG), traitement d'image"),
    ("Temps réel & Natif", "WebRTC, LiveKit, Socket.IO, chiffrement de bout en bout, modules natifs React Native/Expo, monorepos Turborepo/Yarn Berry"),
    ("Cloud & DevOps", "AWS, Google Cloud, Docker, CI/CD, Git"),
    ("Sécurité & Réseaux", "Wireshark, Burp Suite, Nmap, Kali Linux, fondamentaux réseau"),
    ("Outils & Méthodes", "Figma, Postman, QGIS, Odoo, Claude (développement assisté par IA), UML, Agile, Merise (MCD/MLD/MCT), Design Patterns"),
    ("Publication & Déploiement", "Google Play Console, Apple App Store Connect / TestFlight, mises à jour OTA CodePush, Docker, CI/CD GitHub Actions, pipelines de release versionnés"),
    ("Marketing digital", "Gestion de pages Facebook & LinkedIn, création de campagnes Meta Ads Manager, marketing sur les réseaux sociaux"),
]
for label, value in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label} : ")
    r1.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(value)
    r2.font.size = Pt(10.5)

# ---------- EXPERIENCE ----------
add_heading("Expérience Projets (Freelance)")

projects = [
    ("Nougta – Système de Point de Vente Multi-Boutiques", "2025 – Présent", [
        "Développement d'une application POS mobile offline-first (React Native, WatermelonDB) avec synchronisation multi-appareils en temps réel via Firebase Firestore et FCM.",
        "Développement du backend NestJS/MongoDB ainsi que du tableau de bord admin, de la boutique e-commerce et du module comptabilité en Next.js.",
        "Gestion complète des mises en production : builds AAB/APK et iOS versionnés, soumissions Play Store et App Store Connect, mises à jour OTA CodePush, et un script de contrôle de version personnalisé appliquant les règles de versioning Play Store/App Store en CI.",
        "Mise en place du déploiement Docker et des pipelines CI/CD GitHub Actions pour le backend et les applications web.",
    ]),
    ("Plateforme de Communication d'Entreprise (Confidentiel) – Modules Natifs de Chiffrement E2E & WebRTC", "2025 – Présent", [
        "Contribution à des modules natifs React Native/Expo pour une plateforme de communication sécurisée en marque blanche destinée à des clients institutionnels, axée sur les appels vocaux/vidéo et la messagerie chiffrés de bout en bout, en pair-à-pair et en groupe.",
        "Développement de fonctionnalités d'appel WebRTC en temps réel au-dessus de LiveKit et de la signalisation Socket.IO pour les appels P2P et de groupe.",
        "Développement d'une fonctionnalité de rapport journalier permettant à chaque utilisateur de soumettre un rapport d'activité complet, agrégé pour le suivi opérationnel.",
        "Travail au sein d'un monorepo Yarn/Turborepo aux côtés d'équipes dédiées mobile, dashboard et infrastructure.",
    ]),
    ("Khidmaty – Plateforme Nationale de Services Publics (Mobile, Backend & Exploitation)", "Nov 2024 – Présent", [
        "Maintenance complète d'une plateforme nationale de services publics comptant plus de 500 000 utilisateurs, couvrant l'immatriculation/transfert de propriété de véhicules (ANRPTS, carte grise) et les demandes de casier judiciaire avec aperçu, partage et téléchargement de PDF dans l'application.",
        "Administration de la base de données de production, incluant sauvegardes régulières et contrôles d'intégrité des données.",
        "Publication et maintenance de l'application sur Google Play et l'Apple App Store ; résolution des revues de conformité aux politiques des stores pour maintenir les versions approuvées et en ligne.",
        "Diagnostic et correction des bugs remontés par les clients ainsi que des plantages/problèmes de mémoire de l'application mobile ; mise à jour des dépendances et des cibles OS/SDK vers les dernières versions supportées.",
        "Surveillance de la plateforme contre les menaces de sécurité et application de correctifs pour protéger les données des utilisateurs et de l'État.",
        "Représentation de l'équipe technique lors de réunions avec des responsables ministériels et des parties prenantes gouvernementales de haut niveau sur les besoins, incidents et la feuille de route.",
    ]),
    ("Houwity – Application Mobile de Registre National (Frontend)", "Nov 2024 – Présent", [
        "Développement et maintenance du frontend de l'application mobile Houwity, une plateforme de registre national équivalente au portail officiel ANRPTS de Mauritanie (anrpts.gov.mr), donnant aux citoyens un accès mobile aux mêmes services de registre que l'application Khidmaty.",
        "Maintenance continue : mises à jour de fonctionnalités, corrections de bugs et alignement avec les évolutions du backend/API du système de registre.",
    ]),
    ("Temwine – Application Mobile de Chaîne d'Approvisionnement pour Épiceries", "Oct 2024 – Présent", [
        "Développement d'une application React Native avec un backend Next.js/MongoDB pour aider les familles à faible revenu à acheter des produits subventionnés par l'État.",
        "Développement d'API pour vérifier l'éligibilité des bénéficiaires via le registre national d'identité et appliquer une limite d'un achat par jour.",
        "Mise en place de la vérification par OTP et des notifications d'arrivage pour les gérants de magasin.",
    ]),
    ("Plateforme d'Annonces Immobilières", "2024", [
        "Développement d'un site d'annonces immobilières avec filtrage par ville pour la location et la vente.",
        "Amélioration du SEO et mise en place du suivi analytique (TTFB, FID, FCP) avec Matomo et Google Analytics pour suivre le comportement des utilisateurs et le taux de rebond.",
        "Technologies : Next.js, Supabase, Tailwind CSS.",
    ]),
    ("Chatbot IA avec Retrieval-Augmented Generation", "2024", [
        "Développement d'une application de chat basée sur GPT avec streaming des réponses en temps réel via l'OpenAI SDK.",
        "Mise en place du RAG pour répondre aux requêtes spécifiques à l'entreprise et aux requêtes générales.",
        "Technologies : Next.js, Supabase.",
    ]),
    ("Plateforme E-Commerce", "Mars – Avr 2024", [
        "Développement d'un site e-commerce avec navigation produits, panier et paiement, ainsi qu'un tableau de bord admin pour la gestion des produits, des stocks et de l'entrepôt.",
        "Déploiement sur AWS et Hostinger, configuration d'un VPS avec nom de domaine personnalisé pour l'hébergement en production.",
        "Technologies : Spring Boot, Hibernate/JPA, React.",
    ]),
    ("Générateur de Formulaires Personnalisés", "2023 – 2024", [
        "Développement d'un générateur de formulaires par glisser-déposer (upload de fichiers, dates, cases à cocher, boutons radio, champs texte) avec personnalisation du style en direct.",
        "Ajout de l'export du code vers Flutter, React, Next.js et HTML, ainsi que le partage de formulaires et la collecte de soumissions.",
        "Technologies : Next.js, Supabase.",
    ]),
    ("Application Mobile de Réservation de Stades", "Sept – Déc 2023", [
        "Développement d'une application mobile de réservation de stades de football, avec consultation des disponibilités et envoi de preuve de paiement pour validation par le propriétaire.",
        "Développement d'outils côté propriétaire pour gérer les stades et suivre les réservations.",
        "Technologies : Flutter, Dart, Django, MySQL, REST API, GetX, Google Maps.",
    ]),
    ("Système de Point de Vente & Gestion des Stocks", "Juil 2023", [
        "Développement d'une application desktop de point de vente prenant en charge les paiements en espèces, à crédit, avec TVA et par banque, avec alertes de stock bas et rapports de ventes/profits.",
        "Mise en place du suivi des dettes clients, de l'historique des transactions et des sauvegardes automatiques de la base de données.",
        "Technologies : Java, Swing, MySQL, Hibernate, MVC.",
    ]),
    ("Système de Gestion de Flotte de Bus & Passagers", "Juil 2023", [
        "Développement d'une application desktop pour une agence de transport permettant de gérer les trajets inter-villes, la réservation de sièges et le suivi des employés/colis.",
        "Prévention des doublons d'attribution de sièges et génération de manifestes de passagers imprimables par trajet.",
        "Technologies : Python, Tkinter, MySQL.",
    ]),
    ("Plateforme de Freelance", "2023", [
        "Développement d'une marketplace de freelances mettant en relation freelances et clients, avec profils, demandes de mission et perception de commissions.",
        "Technologies : HTML, CSS, JavaScript, PHP, MySQL, Bootstrap.",
    ]),
    ("Système de Gestion des Stocks – Ministère de la Transformation Numérique", "2023", [
        "Développement d'un système de suivi des stocks, des transactions et des rapports pour un ministère.",
        "Technologies : Spring Boot, Hibernate/JPA, Angular.",
    ]),
    ("Site Web de Réservation de Voyages en Bus", "2020", [
        "Développement d'un site de réservation de sièges de bus et de gestion des passagers, avec listes de passagers imprimables pour les contrôles de sécurité.",
        "Renforcement de la sécurité de l'application contre les injections SQL, XSS et HTML.",
        "Technologies : Spring Boot, Hibernate/JPA, React, MySQL.",
    ]),
]

for title, dates, bullets in projects:
    add_entry_title(title, dates)
    for b in bullets:
        add_bullet(b)

# ---------- FORMATION ACADEMIQUE ----------
add_heading("Formation Académique")
add_entry_title("Doctorat en Systèmes Distribués Sécurisés (2ᵉ année) – Université de Nouakchott", "2024 – Présent")
add_entry_title("Master en Technologies de l'Information – Université de Nouakchott", "2022 – 2024")
add_entry_title("Licence en Administration du Développement (Internet & Intranet) – Université de Nouakchott", "2019 – 2022")
add_entry_title("Baccalauréat, Mathématiques (Bac C) – Lycée d'Atar", "2016 – 2019")

# ---------- FORMATIONS PROFESSIONNELLES ----------
add_heading("Formations Professionnelles")
add_bullet("Intelligence Artificielle, Bureautique, Gestion d'Entreprise et Soft Skills (2020 – 2023)")
add_bullet("Marketing Digital : gestion de pages Facebook & LinkedIn et création de campagnes publicitaires (2020 – 2023)")

# ---------- LANGUES ----------
add_heading("Langues")
add_plain("Arabe (Langue maternelle)  •  Français (Courant)  •  Anglais (Courant)", space_after=2)

doc.save("/Users/deidinecheigeur/Downloads/Deidine_Sidina_CV_ATS_FR.docx")
print("saved")
