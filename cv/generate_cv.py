"""
Generate ATS-friendly CVs (DOCX + PDF, EN + FR) from cv_data.json.

Usage:
    .venv/bin/python generate_cv.py
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, ListFlowable, ListItem,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

HERE = Path(__file__).parent
DATA_PATH = HERE / "cv_data.json"
FONT = "Calibri"


def load_data():
    with open(DATA_PATH, encoding="utf-8") as f:
        return json.load(f)


def contact_links(contact: dict):
    """Build (href, display_text) pairs from a structured contact dict."""
    parts = []
    if contact.get("email"):
        parts.append((f"mailto:{contact['email']}", contact["email"]))
    if contact.get("phone"):
        tel = "tel:" + contact["phone"].replace(" ", "")
        parts.append((tel, contact["phone"]))
    if contact.get("linkedin"):
        url = contact["linkedin"]
        href = url if url.startswith("http") else f"https://{url}"
        parts.append((href, url))
    if contact.get("github"):
        url = contact["github"]
        href = url if url.startswith("http") else f"https://{url}"
        parts.append((href, url))
    return parts


# ----------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------
def build_docx(d: dict, out_path: Path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(12)
        pPr = p._p.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pPr.makeelement(
            qn("w:bottom"),
            {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "000000"},
        )
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def add_entry_title(title, meta):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(title)
        r1.bold = True
        r1.font.size = Pt(10.5)
        if meta:
            p.add_run("\t")
            r2 = p.add_run(meta)
            r2.italic = True
            r2.font.size = Pt(10)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.22)
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        return p

    def add_plain(text, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        return p

    def add_hyperlink(paragraph, url, text, size=10):
        part = paragraph.part
        r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)

        new_run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")

        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), FONT)
        rFonts.set(qn("w:hAnsi"), FONT)
        rPr.append(rFonts)

        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size * 2))
        rPr.append(sz)

        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1155CC")
        rPr.append(color)

        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)

        new_run.append(rPr)
        t = OxmlElement("w:t")
        t.text = text
        new_run.append(t)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)
        return hyperlink

    # Header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)
    r = name_p.add_run(d["name"])
    r.bold = True
    r.font.size = Pt(20)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    r = title_p.add_run(d["job_title"])
    r.font.size = Pt(12.5)
    r.italic = True

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(2)
    for i, (href, text) in enumerate(contact_links(d["contact"])):
        if i > 0:
            sep = contact_p.add_run("  |  ")
            sep.font.size = Pt(10)
        add_hyperlink(contact_p, href, text, size=10)

    # Summary
    add_heading(d["summary_heading"])
    add_plain(d["summary_text"], space_after=2)

    # Skills
    add_heading(d["skills_heading"])
    for label, value in d["skills"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}: ")
        r1.bold = True
        r1.font.size = Pt(10.5)
        r2 = p.add_run(value)
        r2.font.size = Pt(10.5)

    # Experience
    add_heading(d["experience_heading"])
    for proj in d["projects"]:
        add_entry_title(proj["title"], proj["dates"])
        for b in proj["bullets"]:
            add_bullet(b)

    # Education
    add_heading(d["education_heading"])
    for edu in d["education"]:
        add_entry_title(edu["title"], edu["dates"])

    # Training
    add_heading(d["training_heading"])
    for t in d["training"]:
        add_bullet(t)

    # Languages
    add_heading(d["languages_heading"])
    add_plain(d["languages_text"], space_after=2)

    doc.save(str(out_path))


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def build_pdf(d: dict, out_path: Path):
    styles = getSampleStyleSheet()

    name_style = ParagraphStyle(
        "Name", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=20, leading=24, alignment=TA_CENTER, spaceAfter=2,
    )
    title_style = ParagraphStyle(
        "JobTitle", parent=styles["Normal"], fontName="Helvetica-Oblique",
        fontSize=12.5, leading=16, alignment=TA_CENTER, spaceAfter=4,
    )
    contact_style = ParagraphStyle(
        "Contact", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10, leading=13, alignment=TA_CENTER, spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "Heading", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=12, leading=15, spaceBefore=12, spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10.5, leading=13.5, spaceAfter=2,
    )
    entry_title_style = ParagraphStyle(
        "EntryTitle", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=10.5, leading=13,
    )
    entry_dates_style = ParagraphStyle(
        "EntryDates", parent=styles["Normal"], fontName="Helvetica-Oblique",
        fontSize=10, leading=13, alignment=2,  # right
    )
    bullet_style = ParagraphStyle(
        "Bullet", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10.5, leading=13.5, leftIndent=14,
    )
    skill_style = ParagraphStyle(
        "Skill", parent=styles["Normal"], fontName="Helvetica",
        fontSize=10.5, leading=13.5, spaceAfter=2,
    )

    story = []
    story.append(Paragraph(d["name"], name_style))
    story.append(Paragraph(d["job_title"], title_style))
    contact_markup = "  |  ".join(
        f'<a href="{href}"><font color="#1155CC">{text}</font></a>'
        for href, text in contact_links(d["contact"])
    )
    story.append(Paragraph(contact_markup, contact_style))

    def add_heading(text):
        story.append(Paragraph(text.upper(), heading_style))
        story.append(HRFlowable(width="100%", thickness=0.75, color="#000000", spaceAfter=4))

    def add_entry(title, dates):
        t = Table(
            [[Paragraph(title, entry_title_style), Paragraph(dates, entry_dates_style)]],
            colWidths=[4.6 * inch, 1.9 * inch],
        )
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        story.append(t)

    def add_bullets(items):
        story.append(ListFlowable(
            [ListItem(Paragraph(b, bullet_style), spaceAfter=2) for b in items],
            bulletType="bullet", start="•", leftIndent=14,
        ))

    # Summary
    add_heading(d["summary_heading"])
    story.append(Paragraph(d["summary_text"], body_style))

    # Skills
    add_heading(d["skills_heading"])
    for label, value in d["skills"]:
        story.append(Paragraph(f"<b>{label}:</b> {value}", skill_style))

    # Experience
    add_heading(d["experience_heading"])
    for proj in d["projects"]:
        add_entry(proj["title"], proj["dates"])
        add_bullets(proj["bullets"])

    # Education
    add_heading(d["education_heading"])
    for edu in d["education"]:
        add_entry(edu["title"], edu["dates"])

    # Training
    add_heading(d["training_heading"])
    add_bullets(d["training"])

    # Languages
    add_heading(d["languages_heading"])
    story.append(Paragraph(d["languages_text"], body_style))

    pdf = SimpleDocTemplate(
        str(out_path), pagesize=LETTER,
        topMargin=0.6 * inch, bottomMargin=0.6 * inch,
        leftMargin=0.75 * inch, rightMargin=0.75 * inch,
        title=f"{d['name']} - {d['job_title']}",
    )
    pdf.build(story)


# ----------------------------------------------------------------------
VARIANT_LABELS = {
    "tech": "Tech",
    "general": "General",
}


def main():
    data = load_data()["variants"]
    out_dir = HERE / "output"
    out_dir.mkdir(exist_ok=True)

    for variant, vlabel in VARIANT_LABELS.items():
        if variant not in data:
            continue
        for lang, label in [("en", "EN"), ("fr", "FR")]:
            d = data[variant][lang]
            docx_path = out_dir / f"Deidine_CV_{vlabel}_{label}.docx"
            pdf_path = out_dir / f"Deidine_CV_{vlabel}_{label}.pdf"
            build_docx(d, docx_path)
            build_pdf(d, pdf_path)
            print(f"[{vlabel}/{label}] {docx_path}")
            print(f"[{vlabel}/{label}] {pdf_path}")


if __name__ == "__main__":
    main()
