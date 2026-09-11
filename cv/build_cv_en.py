from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

FONT = "Calibri"

doc = Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

# Base style
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
rpr = normal.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = rpr.makeelement(qn('w:rFonts'), {})
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), FONT)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)

def add_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.border_id = None
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = None
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pPr.makeelement(qn('w:bottom'), {
        qn('w:val'): 'single', qn('w:sz'): '6', qn('w:space'): '1', qn('w:color'): '000000'
    })
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_entry_title(title, meta):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(10.5)
    if meta:
        tab = p.add_run("\t")
        r2 = p.add_run(meta)
        r2.italic = True
        r2.font.size = Pt(10)
        # right tab stop
        from docx.enum.text import WD_TAB_ALIGNMENT
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    return p

def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    return p

def add_plain(text, bold=False, italic=False, size=10.5, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p

# ---------- HEADER ----------
name_p = doc.add_paragraph()
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
name_p.paragraph_format.space_after = Pt(2)
r = name_p.add_run("DEIDINE SIDINA")
r.bold = True
r.font.size = Pt(20)

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_after = Pt(4)
r = title_p.add_run("Software Engineer")
r.font.size = Pt(12.5)
r.italic = True

contact_p = doc.add_paragraph()
contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact_p.paragraph_format.space_after = Pt(2)
r = contact_p.add_run(
    "cheigeurdeidine@gmail.com  |  +222 49619609  |  linkedin.com/in/deidine  |  github.com/deidine"
)
r.font.size = Pt(10)

# ---------- SUMMARY ----------
add_heading("Professional Summary")
add_plain(
    "Software engineer with experience delivering full-stack web, mobile, and desktop "
    "applications across real estate, e-commerce, government, and fintech-adjacent projects. "
    "Skilled in React/Next.js, React Native, Spring Boot, and Python, with hands-on experience "
    "in AI integration (OpenAI SDK, RAG), cloud deployment (AWS, GCP), and both SQL and NoSQL "
    "database design.",
    space_after=2,
)

# ---------- SKILLS ----------
add_heading("Skills")
skills = [
    ("Languages", "JavaScript, TypeScript, Python, Java, C, C++, PHP, Dart, SQL"),
    ("Frontend & Mobile", "React, Next.js, Angular, Vue.js, React Native, Electron, HTML, CSS, Tailwind CSS, Bootstrap"),
    ("Backend", "Spring Boot, Hibernate/JPA, Laravel, Django, REST API, SOAP"),
    ("Databases", "MySQL, PostgreSQL, MongoDB, Firebase, Supabase, SQL Server, NoSQL"),
    ("AI & Data", "Machine Learning, Deep Learning, NLP, OpenAI SDK, Retrieval-Augmented Generation (RAG), Image Processing"),
    ("Real-Time & Native", "WebRTC, LiveKit, Socket.IO, end-to-end encryption, native React Native/Expo modules, Turborepo/Yarn Berry monorepos"),
    ("Cloud & DevOps", "AWS, Google Cloud, Docker, CI/CD, Git"),
    ("Security & Networking", "Wireshark, Burp Suite, Nmap, Kali Linux, network fundamentals"),
    ("Tools & Methods", "Figma, Postman, QGIS, Odoo, Claude (AI-assisted development), UML, Agile, Merise (MCD/MLD/MCT), Design Patterns"),
    ("Release & Deployment", "Google Play Console, Apple App Store Connect / TestFlight, CodePush OTA updates, Docker, GitHub Actions CI/CD, Fastlane-style versioned release pipelines"),
    ("Digital Marketing", "Facebook & LinkedIn Page management, Meta Ads Manager campaign creation, social media marketing"),
]
for label, value in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{label}: ")
    r1.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(value)
    r2.font.size = Pt(10.5)

# ---------- EXPERIENCE ----------
add_heading("Project Experience (Freelance)")

projects = [
    ("Nougta – Multi-Tenant Point-of-Sale System", "2025 – Present", [
        "Building an offline-first POS mobile app (React Native, WatermelonDB) with real-time multi-device sync via Firebase Firestore and FCM.",
        "Developed the NestJS/MongoDB backend and Next.js admin dashboard, ecommerce storefront, and accounting module.",
        "Own end-to-end release management: versioned AAB/APK and iOS release builds, Play Store and App Store Connect submissions, CodePush OTA updates, and a custom version-guard script enforcing Play Store/App Store versioning rules in CI.",
        "Set up Docker-based deployment and GitHub Actions CI/CD pipelines for the backend and web apps.",
    ]),
    ("Enterprise Communication Platform (Confidential) – Native E2E Encryption & WebRTC Modules", "2025 – Present", [
        "Contributed native React Native/Expo modules to a whitelabel, multi-client secure communication platform (institutional clients), focused on end-to-end encrypted peer-to-peer and group voice/video calling and messaging.",
        "Built real-time WebRTC calling features on top of LiveKit and Socket.IO signaling for P2P and group calls.",
        "Built a daily reporting feature letting each user submit a full activity report, aggregated for operational review.",
        "Worked within a Yarn/Turborepo monorepo alongside dedicated mobile, dashboard, and infrastructure teams.",
    ]),
    ("Khidmaty – National Government Services Platform (Mobile, Backend & Ops)", "Nov 2024 – Present", [
        "Own end-to-end maintenance of a national government services platform serving 500,000+ users, covering vehicle registration/ownership transfer (ANRPTS, carte grise) and judicial record (casier judiciaire) requests with in-app PDF preview, share, and download.",
        "Administer the production database, including regular backups and data integrity checks.",
        "Release and maintain the app on both Google Play and the Apple App Store; resolve store policy reviews and compliance issues to keep releases approved and live.",
        "Triage and fix client-reported bugs and mobile app crashes/memory issues; upgrade dependencies and OS/SDK targets to the latest supported versions.",
        "Monitor the platform for security threats and apply security patches to protect user and government data.",
        "Represent the engineering side in meetings with ministry officials and senior government stakeholders on requirements, incidents, and roadmap.",
    ]),
    ("Houwity – National Registry Mobile App (Frontend)", "Nov 2024 – Present", [
        "Built and maintain the mobile app frontend for Houwity, a national registry platform equivalent to Mauritania's official ANRPTS portal (anrpts.gov.mr), giving citizens mobile access to the same registry services as the Khidmaty app.",
        "Own ongoing maintenance: feature updates, bug fixes, and alignment with backend/API changes as the registry system evolves.",
    ]),
    ("Temwine – Grocery Supply Chain Mobile App", "Oct 2024 – Present", [
        "Developing a React Native app with a Next.js/MongoDB backend to help low-income families purchase government-subsidized groceries.",
        "Built APIs to verify beneficiary eligibility against the national ID registry and enforce one-purchase-per-day limits.",
        "Implemented OTP verification and shipment notifications for store owners.",
    ]),
    ("Real Estate Listing Platform", "2024", [
        "Built a property listing website with city-based filtering for rentals and sales.",
        "Improved SEO and set up analytics tracking (TTFB, FID, FCP) with Matomo and Google Analytics to monitor user behavior and bounce rate.",
        "Technologies: Next.js, Supabase, Tailwind CSS.",
    ]),
    ("AI Chatbot with Retrieval-Augmented Generation", "2024", [
        "Built a GPT-based chat application with real-time response streaming using the OpenAI SDK.",
        "Implemented RAG to answer company-specific and general-knowledge queries.",
        "Technologies: Next.js, Supabase.",
    ]),
    ("E-Commerce Platform", "Mar – Apr 2024", [
        "Built an e-commerce site with product browsing, cart, and checkout, plus an admin dashboard for product, inventory, and warehouse management.",
        "Deployed to AWS and Hostinger and configured a VPS with a custom domain for production hosting.",
        "Technologies: Spring Boot, Hibernate/JPA, React.",
    ]),
    ("Custom Form Builder", "2023 – 2024", [
        "Built a drag-and-drop form builder (file uploads, dates, checkboxes, radio buttons, text fields) with live style customization.",
        "Added code export to Flutter, React, Next.js, and HTML, plus form sharing and submission collection.",
        "Technologies: Next.js, Supabase.",
    ]),
    ("Stadium Booking Mobile App", "Sep – Dec 2023", [
        "Built a mobile app for booking football stadiums, with availability browsing and payment-proof upload for owner validation.",
        "Built owner-side tools for managing stadiums and tracking reservations.",
        "Technologies: Flutter, Dart, Django, MySQL, REST API, GetX, Google Maps.",
    ]),
    ("Point-of-Sale & Inventory Management System", "Jul 2023", [
        "Built a desktop POS application supporting cash, credit, VAT, and bank payment types, with low-stock alerts and sales/profit reporting.",
        "Implemented client debt tracking, transaction history, and automated database backups.",
        "Technologies: Java, Swing, MySQL, Hibernate, MVC.",
    ]),
    ("Bus Fleet & Passenger Management System", "Jul 2023", [
        "Built a desktop app for a bus agency to manage inter-city trips, seat booking, and worker/package tracking.",
        "Prevented duplicate seat assignments and generated printable passenger manifests per trip.",
        "Technologies: Python, Tkinter, MySQL.",
    ]),
    ("Freelance Marketplace Platform", "2023", [
        "Built a freelance marketplace connecting freelancers with clients, including profiles, job requests, and fee collection.",
        "Technologies: HTML, CSS, JavaScript, PHP, MySQL, Bootstrap.",
    ]),
    ("Inventory Management System – Ministry of Technology Transformation", "2023", [
        "Developed an inventory tracking system for stock levels, transactions, and reporting for a government ministry.",
        "Technologies: Spring Boot, Hibernate/JPA, Angular.",
    ]),
    ("Bus Travel Booking Website", "2020", [
        "Built a website for bus seat booking and passenger management, with printable passenger lists for security checkpoints.",
        "Hardened the application against SQL injection, XSS, and HTML injection.",
        "Technologies: Spring Boot, Hibernate/JPA, React, MySQL.",
    ]),
]

for title, dates, bullets in projects:
    add_entry_title(title, dates)
    for b in bullets:
        add_bullet(b)

# ---------- EDUCATION ----------
add_heading("Education")
add_entry_title("Ph.D. in Secure Distributed Systems (2nd year) – University of Nouakchott", "2024 – Present")
add_entry_title("Master's in Information Technology – University of Nouakchott", "2022 – 2024")
add_entry_title("Bachelor's in Development Administration (Internet & Intranet) – University of Nouakchott", "2019 – 2022")
add_entry_title("Baccalauréat, Mathematics (Bac C) – Lycée Atar", "2016 – 2019")

# ---------- TRAINING ----------
add_heading("Professional Training")
add_bullet("Artificial Intelligence, Office Automation, Business Management, and Soft Skills (2020 – 2023)")
add_bullet("Digital Marketing: Facebook & LinkedIn Page management and paid ad campaign creation (2020 – 2023)")

# ---------- LANGUAGES ----------
add_heading("Languages")
add_plain("Arabic (Native)  •  French (Fluent)  •  English (Fluent)", space_after=2)

doc.save("/Users/deidinecheigeur/Downloads/Deidine_Sidina_CV_ATS.docx")
print("saved")
